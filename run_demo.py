import os
import argparse
import math
import subprocess
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path
import openpyxl
import json
from communication_cost import*

parser = argparse.ArgumentParser()
parser.add_argument("-p", default=None, help="existing instruction trace root; otherwise generate fresh traces")
parser.add_argument("-n", default="Sheet", help="sheet name")
parser.add_argument("-m", default="OPT-2.7B", help="Model to run")
parser.add_argument("-i", default=1024, help="input tokens")
parser.add_argument("-o", default=128, help="output tokens")
parser.add_argument("-b", default=64, help="batch size")
parser.add_argument("-s", default="configs/HBNPU_8Gb.ini", help="specification of DRAM")


@contextmanager
def simulate_kernel(trace, spec, log, *, binary="./build/dramsim3main", cycles=10000000):
    """Require END completion; never read stale output from a previous kernel."""
    with tempfile.TemporaryDirectory(prefix="hbnpu-opt-kernel-") as output:
        with open(log, "a") as stream:
            stream.write(str(trace) + "\n")
            stream.flush()
            subprocess.run([str(binary), str(spec), "-c", str(cycles),
                            "-t", str(trace), "-o", output],
                           stdout=stream, stderr=subprocess.STDOUT, check=True)
        metrics = json.loads((Path(output) / "hbnpu_instruction.json").read_text())
        if not metrics["end_completed"]:
            raise RuntimeError("instruction program did not complete: " + str(trace))
        yield json.loads((Path(output) / "dramsim3.json").read_text())

def run(path, name, model, in_tokens, out_tokens, batch_size, spec):

    if not os.path.exists('logs'):
        os.makedirs('logs')

    print("Trace file checking..")
    workload = "_".join([model, str(in_tokens), str(out_tokens), str(batch_size)])
    print("Workload: " + workload)
    trace_folder = os.path.join(path, workload)
    if not os.path.exists(trace_folder):
        raise Exception("Traces do not exist! Generate traces first.")

    print("Model configuration checking..")
    fin_ = open('models', 'r')
    lines = fin_.readlines()
    lines_parsed = [line.strip().split(' ') for line in lines if line.strip().split(' ')[0] == model]
    if not lines_parsed:
        raise Exception("Model Not Found!")
    else:
        _, params, n_layers, d_model, n_heads, d_head, TP, PP = lines_parsed[0]


    layer_dict = {"createQKV":'B', "QK":'C', "SV":'D', "Wo":'E', "L1":'F', "L2":'G'}
    if os.path.exists('result.xlsx'):
        wb = openpyxl.load_workbook('result.xlsx')
    else:
        wb = openpyxl.Workbook()

    if workload in wb.sheetnames:
        wb.remove(wb[workload])

    wb.create_sheet(title=workload)
    sh = wb[workload]

    # TODO assuming tCK=1ns
    sh['A'+str(2)] = 'prefill runtime (ns)'
    sh['A'+str(3)] = 'prefill energy (pJ)'
    sh['A'+str(4)] = 'decode runtime (ns)'
    sh['A'+str(5)] = 'decode energy (pJ)'

    sh['B'+str(1)] = 'createQKV'
    sh['C'+str(1)] = 'QK'
    sh['D'+str(1)] = 'SV'
    sh['E'+str(1)] = 'Wo'
    sh['F'+str(1)] = 'L1'
    sh['G'+str(1)] = 'L2'
    sh['H'+str(1)] = 'total'
    sh['I'+str(1)] = 'throughput (K tokens/s)'

    print("\nRunning prefill (summarization) phase..")
    prompt_folder = trace_folder + "/prompt/"
    flist = sorted(os.listdir(prompt_folder))
    for f in flist:

        with simulate_kernel(prompt_folder + f, spec, "logs/" + workload + "_prompt.log") as json_object:
            cycles = json_object['0']['num_cycles']
            energy = sum([json_object[str(ch)]['total_energy'] for ch in range(0,8)])
            if cycles == 10000000:
                raise Exception("compute not finished! prompt: " + f)
            cycles *= int(n_layers)*int(max(batch_size/int(PP),1))
            energy *= int(n_layers)*max(batch_size/int(PP),1)
            if f == "createQKV":
                cycles *= 3
                energy *= 3
            elif f == "QK" or f == "SV":
                cycles *= int(max(int(math.ceil(float(n_heads)/int(TP))), 1))
                energy *= max(int(math.ceil(float(n_heads)/int(TP))), 1)


            sh[layer_dict[f] + str(2)] = cycles
            sh[layer_dict[f] + str(3)] = energy
    total_cycles = sum([int(str(sh[a+str(2)].value)) for a in ['B', 'C', 'D', 'E', 'F', 'G'] if sh[a+str(2)].value is not None])
    total_energy = sum([float(str(sh[a+str(3)].value)) for a in ['B', 'C', 'D', 'E', 'F', 'G'] if sh[a+str(3)].value is not None])
    chips_per_board = 16 # 32 GB per board
    communication_latency = cost_communication(int(n_layers), int(d_model), in_tokens, batch_size, chips_per_board, int(TP), int(PP), False)
    vector_latency = vector_process(int(n_layers), int(d_model), int(n_heads), int(d_head), int(TP), int(PP), in_tokens, out_tokens, batch_size, False)
    total_cycles = max(total_cycles + vector_latency, communication_latency)
    sh['H'+str(2)] = total_cycles
    sh['H'+str(3)] = total_energy


    print("Prefill runtime: " + str(total_cycles) + " ns")
    print("Prefill energy: {:.2f} pJ".format(total_energy))


    print("\nRunning decode (generation) phase..")

    decode_folder = trace_folder + "/decode/"
    flist = sorted(os.listdir(decode_folder))

    QK_cycles, QK_energy, SV_cycles, SV_energy = 0, 0, 0, 0
    for f in flist:
        if os.path.isdir(decode_folder + f):
            for ff in os.listdir(decode_folder + f):

                with simulate_kernel(decode_folder + f + "/" + ff, spec, "logs/" + workload + "_decode.log") as json_object:
                    cycles = json_object['0']['num_cycles']*int(n_layers)
                    energy = sum([json_object[str(ch)]['total_energy'] for ch in range(0,8)])*int(n_layers)
                    if cycles == 10000000:
                        raise Exception("compute not finished! decode: "+ f +' '+ ff)
                    if "QK_" in ff:
                        # parallelized by 16 Subarrays
                        QK_cycles += cycles*int(max(int(math.ceil(float(n_heads)/int(TP)))*batch_size/int(PP)/16, 1))
                        QK_energy += energy*max(int(math.ceil(float(n_heads)/int(TP)))*batch_size/int(PP)/16, 1)
                    elif "SV_" in ff:
                        SV_cycles += cycles*int(max(int(math.ceil(float(n_heads)/int(TP)))*batch_size/int(PP)/16, 1))
                        SV_energy += energy*max(int(math.ceil(float(n_heads)/int(TP)))*batch_size/int(PP)/16, 1)
                    else:
                        cycles *= out_tokens
                        energy *= out_tokens
                        if ff == "createQKV":
                            cycles *= 3
                            energy *= 3
                        sh[layer_dict[ff] + '4'] = int(min(cycles, int(str(sh[layer_dict[ff]+'4'].value)))) if sh[layer_dict[ff]+'4'].value is not None else cycles
                        sh[layer_dict[ff] + '5'] = min(energy, float(str(sh[layer_dict[ff]+'5'].value))) if sh[layer_dict[ff]+'5'].value is not None else energy



        else:

            with simulate_kernel(decode_folder + f, spec, "logs/" + workload + "_decode.log") as json_object:
                cycles = json_object['0']['num_cycles']*int(n_layers)
                energy = sum([json_object[str(ch)]['total_energy'] for ch in range(0,8)])*int(n_layers)
                if cycles == 10000000:
                    raise Exception("compute not finished! decode TSGEMM: " + f)
                # granularity of batch size 8 in TS-GEMM dataflow
                cycles *= int(out_tokens*max(batch_size/int(PP)/8,1))
                energy *= out_tokens*max(batch_size/int(PP)/8,1)
                if f == "createQKV":
                    cycles *= 3
                    energy *= 3

                # Compare to WS dataflow
                sh[layer_dict[f] + '4'] = int(min(cycles, int(str(sh[layer_dict[f]+'4'].value)))) if sh[layer_dict[f]+'4'].value is not None else cycles
                sh[layer_dict[f] + '5'] = min(energy, float(str(sh[layer_dict[f]+'5'].value))) if sh[layer_dict[f]+'5'].value is not None else energy

    sh['C' + '4'] = QK_cycles
    sh['C' + '5'] = QK_energy
    sh['D' + '4'] = SV_cycles
    sh['D' + '5'] = SV_energy
    chips_per_board = 16 # 32 GB per board
    communication_latency = cost_communication(int(n_layers), int(d_model), out_tokens, batch_size, chips_per_board, int(TP), int(PP), True)
    vector_latency = vector_process(int(n_layers), int(d_model), int(n_heads), int(d_head), int(TP), int(PP), in_tokens, out_tokens, batch_size, True)

    total_cycles = sum([int(str(sh[a+str(4)].value)) for a in ['B', 'C', 'D', 'E', 'F', 'G'] if sh[a+str(4)].value is not None])
    total_energy = sum([float(str(sh[a+str(5)].value)) for a in ['B', 'C', 'D', 'E', 'F', 'G'] if sh[a+str(5)].value is not None])
    total_cycles += communication_latency + vector_latency
    sh['H'+str(4)] = total_cycles
    sh['H'+str(5)] = total_energy


    throughput = batch_size*out_tokens / total_cycles * 10**6
    sh['I'+str(4)] = throughput

    print("Throughput: {:.2f} K tokens/s".format(throughput))
    print("Energy: {:.2f} pJ".format(total_energy) + '\n')

    wb.save('result.xlsx')
    return

def vector_process(n_layers, d_model, n_heads, d_head, TP, PP, in_tokens, out_tokens, batch_size, GEN):
    VPU_lanes = 128
    logic_freq = 1 #GHz
    if GEN:
        softmax_flops = int(5*n_heads*(2*in_tokens+out_tokens)/2)
    else:
        softmax_flops = int(5*n_heads*in_tokens)
    relu_flops = d_model*4
    norm_flops = 5*d_model
    if GEN:
        tokens = out_tokens
    else:
        tokens = in_tokens
    total_flops = n_layers * batch_size * tokens * (softmax_flops + relu_flops + 2*norm_flops) / TP / PP
    latency = total_flops / VPU_lanes / logic_freq
    return latency





if __name__ == "__main__":
    args = parser.parse_args()
    if args.p is not None:
        run(args.p, args.n, args.m, int(args.i), int(args.o), int(args.b), args.s)
    else:
        # Generate current instruction traces without overwriting supplied paper traces.
        with tempfile.TemporaryDirectory(prefix="hbnpu-opt-") as trace_root:
            subprocess.run([
                sys.executable, str(Path(__file__).with_name("gen_LLM_trace.py")),
                "-p", trace_root, "-m", args.m, "-i", str(args.i),
                "-o", str(args.o), "-b", str(args.b), "-s", args.s,
            ], check=True)
            run(trace_root, args.n, args.m, int(args.i), int(args.o), int(args.b), args.s)
